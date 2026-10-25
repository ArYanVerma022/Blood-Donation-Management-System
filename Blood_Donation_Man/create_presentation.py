import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_jeevanrakt_presentation():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    DARK_BLUE = RGBColor(11, 37, 69)      # #0b2545
    NAVY_BLUE = RGBColor(19, 60, 85)      # #133c55
    BLOOD_RED = RGBColor(220, 53, 69)     # #dc3545
    LIGHT_BG  = RGBColor(240, 247, 255)   # #f0f7ff
    WHITE     = RGBColor(255, 255, 255)
    TEXT_DARK = RGBColor(33, 37, 41)
    TEXT_MUTED = RGBColor(108, 117, 125)
    
    FONT_FAMILY = "Times New Roman"

    blank_slide_layout = prs.slide_layouts[6]

    # --- Helper: Add Header Banner ---
    def add_header(slide, title_text, category_text="JEEVANRAKT PRESENTATION"):
        # Header background bar
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = DARK_BLUE
        header_box.line.color.rgb = DARK_BLUE
        
        # Red accent bar
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = BLOOD_RED
        accent.line.color.rgb = BLOOD_RED

        # Header Text
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(12), Inches(0.9))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p_sub = tf.paragraphs[0]
        p_sub.text = category_text.upper()
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = BLOOD_RED
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # --- Helper: Add Footer ---
    def add_footer(slide, current_slide, total_slides=7):
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.133), Inches(0.4))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"JeevanRakt – Blood Donor Management System  |  Developed by Aryan Verma (241306064)  |  Slide {current_slide} of {total_slides}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    
    # Background fill
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = LIGHT_BG
    bg1.line.color.rgb = LIGHT_BG

    # Hero Banner Box
    hero = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    hero.fill.solid()
    hero.fill.fore_color.rgb = DARK_BLUE
    hero.line.color.rgb = BLOOD_RED

    tf1 = hero.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.6)
    tf1.margin_top = Inches(0.6)
    tf1.margin_right = Inches(0.6)

    p = tf1.paragraphs[0]
    p.text = "🩸 JeevanRakt"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = BLOOD_RED
    p.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "Blood Donor Management System"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "A College-Level BCA Academic Project Presentation"
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(16)
    p3.font.italic = True
    p3.font.color.rgb = RGBColor(208, 225, 249)
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(15)

    # Line Separator
    p_sep = tf1.add_paragraph()
    p_sep.text = "________________________________________________________"
    p_sep.font.color.rgb = BLOOD_RED
    p_sep.alignment = PP_ALIGN.CENTER
    p_sep.space_before = Pt(10)

    p4 = tf1.add_paragraph()
    p4.text = "Project Created By: Aryan Verma"
    p4.font.name = FONT_FAMILY
    p4.font.size = Pt(22)
    p4.font.bold = True
    p4.font.color.rgb = WHITE
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(20)

    p5 = tf1.add_paragraph()
    p5.text = "Student ID: 241306064  |  Technology: PHP + MySQL + Bootstrap 5"
    p5.font.name = FONT_FAMILY
    p5.font.size = Pt(15)
    p5.font.color.rgb = RGBColor(255, 193, 7)
    p5.alignment = PP_ALIGN.CENTER
    p5.space_before = Pt(8)


    # ==========================================
    # SLIDE 2: TABLE OF CONTENTS
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Table of Contents")
    add_footer(slide2, 2)

    topics = [
        ("01. Introduction & Need", "Background, objective, and purpose of JeevanRakt system."),
        ("02. Real-World Applications", "Healthcare use cases, patient emergency support & donor safety."),
        ("03. Technology Stack Used", "PHP (PDO), MySQL Database, HTML5, CSS3, Bootstrap 5 & XAMPP."),
        ("04. Key System Features", "Donor registration, availability search, blood requests & reports."),
        ("05. Pre-Donation Screening", "Automated health questionnaire logic delivering 🟢 🟡 🔴 preliminary status."),
        ("06. Admin Controls & Inventory", "Real-time stock auto-sync, request fulfillment & dashboard metrics."),
        ("07. Conclusion & Roadmap", "Summary of achievements, educational scope, and future enhancements.")
    ]

    left_m = Inches(0.8)
    top_m = Inches(1.5)
    width = Inches(11.733)
    
    for i, (title, desc) in enumerate(topics):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_m, top_m + Inches(i * 0.72), width, Inches(0.62))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = RGBColor(208, 225, 249)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.08)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = TEXT_DARK


    # ==========================================
    # SLIDE 3: INTRODUCTION
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Introduction & Project Background")
    add_footer(slide3, 3)

    # Box 1: Problem Statement
    box1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    box1.fill.solid()
    box1.fill.fore_color.rgb = LIGHT_BG
    box1.line.color.rgb = BLOOD_RED
    tf1 = box1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_top = Inches(0.3)

    p = tf1.paragraphs[0]
    p.text = "📌 Background & Problem Statement"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLOOD_RED

    points_prob = [
        "In critical medical emergencies, finding compatible blood donors quickly is a major challenge.",
        "Traditional paper-based record-keeping leads to data delays, loss of donor history, and inaccurate stock counts.",
        "Lack of donor health screening before donation leads to wasted time and safety risks at blood banks.",
        "Need for a simple, college-level, reliable web platform for small health centers and donor communities."
    ]
    for pt in points_prob:
        p_pt = tf1.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.name = FONT_FAMILY
        p_pt.font.size = Pt(14)
        p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(12)

    # Box 2: System Objectives
    box2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    box2.fill.solid()
    box2.fill.fore_color.rgb = LIGHT_BG
    box2.line.color.rgb = DARK_BLUE
    tf2 = box2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = "🎯 JeevanRakt Core Objectives"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    points_obj = [
        "Digitalize blood donor registration with contact details, blood groups, and city location filtering.",
        "Provide an automated Pre-Donation Health Questionnaire for preliminary safety evaluation.",
        "Maintain real-time inventory tracking for all 8 blood groups (A+, A-, B+, B-, AB+, AB-, O+, O-).",
        "Enable seamless patient blood request submission and automated administrative stock auto-sync."
    ]
    for pt in points_obj:
        p_pt = tf2.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.name = FONT_FAMILY
        p_pt.font.size = Pt(14)
        p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(12)


    # ==========================================
    # SLIDE 4: APPLICATION OF PROJECT
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Real-World Applications of JeevanRakt")
    add_footer(slide4, 4)

    apps = [
        ("🏥 Small Hospitals & Clinics", "Helps local healthcare facilities manage their internal blood units, monitor low-stock thresholds, and fulfill emergency patient requests efficiently without complex commercial software."),
        ("🤝 Community Donor Matchmaking", "Connects individuals seeking urgent blood transfusions with nearby registered voluntary donors based on blood group, city, and active availability status."),
        ("🧑‍⚕️ Preliminary Medical Screening", "Saves medical staff time by collecting donor symptoms, medical history, alcohol intake, and recent procedures prior to formal blood-bank laboratory evaluation."),
        ("📊 Administrative Record Keeping", "Provides clean database logging for donation histories, screening logs, donor status changes, and printable administrative summary reports for college evaluations.")
    ]

    positions = [
        (Inches(0.8), Inches(1.6)),
        (Inches(6.9), Inches(1.6)),
        (Inches(0.8), Inches(4.3)),
        (Inches(6.9), Inches(4.3))
    ]

    for i, (title, text) in enumerate(apps):
        pos_x, pos_y = positions[i]
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, pos_y, Inches(5.6), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = NAVY_BLUE

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = BLOOD_RED

        p_d = tf.add_paragraph()
        p_d.text = text
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_DARK
        p_d.space_before = Pt(8)


    # ==========================================
    # SLIDE 5: TECHNOLOGY USED
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Technology Stack & Architecture")
    add_footer(slide5, 5)

    techs = [
        ("🐘 PHP 8.x (PDO Engine)", "Backend Logic & Security", "Handles server-side data processing, session state authentication, SQL injection prevention via PDO prepared statements, and screening rules logic."),
        ("🐬 MySQL / MariaDB Database", "Relational Storage (7 Tables)", "Stores admin credentials, user accounts, donor profiles, health screening records, donation logs, blood stock units, and patient requests."),
        ("🎨 HTML5, CSS3 & Bootstrap 5", "Responsive UI Design", "Clean, student-project medical visual design with light blue containers, dark blue navbar, blood red buttons, custom SVG graphics, and responsive cards."),
        ("⚙️ XAMPP Server Environment", "Local Development Stack", "Cross-platform local web server environment running Apache HTTP Server and MariaDB MySQL database for local deployment.")
    ]

    for i, (name, tag, desc) in enumerate(techs):
        top_y = Inches(1.5 + i * 1.3)
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_y, Inches(11.733), Inches(1.15))
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT_BG
        box.line.color.rgb = RGBColor(208, 225, 249)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = name + "  |  " + tag
        p.font.name = FONT_FAMILY
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_DARK
        p_d.space_before = Pt(4)


    # ==========================================
    # SLIDE 6: FEATURES INCLUDED
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Key System Features & Screening Logic")
    add_footer(slide6, 6)

    # Left Column: Public Features
    box_l = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = WHITE
    box_l.line.color.rgb = BLOOD_RED
    tfl = box_l.text_frame
    tfl.word_wrap = True
    tfl.margin_left = Inches(0.3)
    tfl.margin_top = Inches(0.2)

    p = tfl.paragraphs[0]
    p.text = "🌐 Public & Donor Modules"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLOOD_RED

    feats_pub = [
        "Donor Registration: Validates age (18-65), blood group, contact info & duplicate checks.",
        "Pre-Donation Health Screening: Evaluates symptoms, medical history, procedures & alcohol.",
        "Eligibility Outcomes: Displays 🟢 Proceed, 🟡 Review Required, or 🔴 Temporarily Deferred.",
        "Donor Directory Search: Search donors by Blood Group, City, and Availability Status.",
        "Patient Blood Request: Submit emergency blood requests directly to the system."
    ]
    for ft in feats_pub:
        p_ft = tfl.add_paragraph()
        p_ft.text = "• " + ft
        p_ft.font.name = FONT_FAMILY
        p_ft.font.size = Pt(13)
        p_ft.font.color.rgb = TEXT_DARK
        p_ft.space_before = Pt(8)

    # Right Column: Admin Features
    box_r = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = WHITE
    box_r.line.color.rgb = DARK_BLUE
    tfr = box_r.text_frame
    tfr.word_wrap = True
    tfr.margin_left = Inches(0.3)
    tfr.margin_top = Inches(0.2)

    p = tfr.paragraphs[0]
    p.text = "🛡️ Admin Control & Inventory Sync"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    feats_adm = [
        "Admin Dashboard: Displays 4 main stat cards, recent donors, and low-stock alerts.",
        "Donation Recording: Logs new donations, auto-increments stock & updates last donation date.",
        "Blood Request Processing: Approves requests and automatically deducts stock upon completion.",
        "Inventory Stock Alerts: Displays low-stock warnings when blood units fall below 5 units.",
        "Printable Summary Reports: Generates a blood group distribution matrix for official reporting."
    ]
    for ft in feats_adm:
        p_ft = tfr.add_paragraph()
        p_ft.text = "• " + ft
        p_ft.font.name = FONT_FAMILY
        p_ft.font.size = Pt(13)
        p_ft.font.color.rgb = TEXT_DARK
        p_ft.space_before = Pt(8)


    # ==========================================
    # SLIDE 7: CONCLUSION
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Conclusion & Future Roadmap")
    add_footer(slide7, 7)

    # Box 1: Conclusion
    c_box1 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    c_box1.fill.solid()
    c_box1.fill.fore_color.rgb = LIGHT_BG
    c_box1.line.color.rgb = DARK_BLUE
    tfc1 = c_box1.text_frame
    tfc1.word_wrap = True
    tfc1.margin_left = Inches(0.3)
    tfc1.margin_top = Inches(0.3)

    p = tfc1.paragraphs[0]
    p.text = "🏆 Project Achievements"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    conc_pts = [
        "Successfully developed JeevanRakt as a functional BCA college web application.",
        "Maintained a clean visual aesthetic (Simple + Functional + Medical Theme) adhering to student project standards.",
        "Implemented robust database operations with PDO prepared statements and security checks.",
        "Demonstrated automated pre-donation health screening and inventory auto-synchronization."
    ]
    for pt in conc_pts:
        p_pt = tfc1.add_paragraph()
        p_pt.text = "✔ " + pt
        p_pt.font.name = FONT_FAMILY
        p_pt.font.size = Pt(14)
        p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(12)

    # Box 2: Future Roadmap
    c_box2 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    c_box2.fill.solid()
    c_box2.fill.fore_color.rgb = LIGHT_BG
    c_box2.line.color.rgb = BLOOD_RED
    tfc2 = c_box2.text_frame
    tfc2.word_wrap = True
    tfc2.margin_left = Inches(0.3)
    tfc2.margin_top = Inches(0.3)

    p = tfc2.paragraphs[0]
    p.text = "🚀 Future Enhancement Scope"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLOOD_RED

    future_pts = [
        "SMS & Email Notifications: Automated alerts to donors when their blood group is urgently needed.",
        "GPS Location Distance Matching: Calculate real-time distance between donors and emergency patient hospitals.",
        "Mobile App Integration: Android app development using Flutter or React Native for portable donor access.",
        "Multi-Hospital Network: Expanding from single blood bank to regional hospital inventory sharing."
    ]
    for pt in future_pts:
        p_pt = tfc2.add_paragraph()
        p_pt.text = "🚀 " + pt
        p_pt.font.name = FONT_FAMILY
        p_pt.font.size = Pt(14)
        p_pt.font.color.rgb = TEXT_DARK
        p_pt.space_before = Pt(12)


    # Save presentation
    output_filename = "JeevanRakt_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'!")

if __name__ == "__main__":
    build_jeevanrakt_presentation()
